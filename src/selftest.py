#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""端到端自检脚本：生成无文本层 PDF -> 跑 OCR -> 校验可搜索双层 PDF。"""
import os, sys, shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "src"))

import fitz
from PIL import Image, ImageDraw, ImageFont

WORK = ROOT / "testdata"
SHOTS = WORK / "shots"


def make_pdf(path: Path, pages: int = 3) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = fitz.open()
    font_path = ""
    for cand in ("/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
                 "/System/Library/Fonts/STHeiti Medium.ttc",
                 "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
                 "C:/Windows/Fonts/msyh.ttc"):
        if os.path.exists(cand):
            font_path = cand
            break
    lines_pool = [
        ["档案数据治理工具", "OCR 全文检索验证", "编号：DA/T 18-2022-0001", "本文件为测试文档，不含文本层。"],
        ["Batch Processing", "File: sample_b.pdf", "Page size: A4", "This page is image-only as well."],
        ["归档整理报告", "生成日期：2026-10-08", "总页数：3", "状态：待归档"],
    ]
    for i in range(pages):
        page = doc.new_page(width=595, height=842)  # A4 pt
        W, H = 1240, 1754  # 约 150dpi
        img = Image.new("RGB", (W, H), "white")
        d = ImageDraw.Draw(img)
        try:
            f_title = ImageFont.truetype(font_path, 56)
            f_body = ImageFont.truetype(font_path, 40)
        except Exception:
            f_title = f_body = ImageFont.load_default()
        d.rectangle([0, 0, W - 1, H - 1], outline="black", width=3)
        d.text((100, 120), lines_pool[i][0], fill="black", font=f_title)
        d.line([100, 210, W - 100, 210], fill="black", width=3)
        for k, ln in enumerate(lines_pool[i][1:]):
            d.text((100, 280 + k * 90), ln, fill="black", font=f_body)
        raw = WORK / f"page{i}.jpg"
        img.save(raw, "JPEG", quality=90)
        page.insert_image(fitz.Rect(0, 0, 595, 842), filename=str(raw))
        raw.unlink(missing_ok=True)
    doc.save(str(path))
    doc.close()


def main() -> int:
    WORK.mkdir(parents=True, exist_ok=True)
    sample = WORK / "sample_a.pdf"
    make_pdf(sample, pages=3)
    print(f"[1] 生成测试 PDF: {sample} ({sample.stat().st_size} bytes)")

    d = fitz.open(str(sample))
    assert d[0].get_text().strip() == "", "测试 PDF 不应自带文本层"
    d.close()
    print("[2] 确认原 PDF 无文本层 ✓")

    import pdf_ocr_desktop as mod

    cfg = dict(mod.DEFAULTS)
    cfg["dpi"] = 300
    cfg["log_preview_lines"] = 5

    logs: list[str] = []
    def log(level, text):
        line = f"[{level}] {text}"
        logs.append(line)
        print("   ", line[:150])

    def progress(cur, total, text):
        print(f"    → {text}")

    def done(records, ok):
        print(f"[6] 任务结束 ok={ok} 记录数={len(records)}")

    from pathlib import Path as P
    task = mod.BatchTask(cfg=cfg, files=[], root=WORK, overwrite=False,
                         out_dir=WORK / "out", report_path=WORK / "out" / mod.REPORT_FILENAME,
                         msg=log, progress=progress, done=done)
    task.run()

    out_pdf = WORK / "out" / "sample_a_ocr.pdf"
    assert out_pdf.exists(), f"未生成输出文件: {out_pdf}"
    print(f"[3] 输出文件已生成: {out_pdf} ({out_pdf.stat().st_size} bytes)")

    d = fitz.open(str(out_pdf))
    pages = d.page_count
    texts = [d[i].get_text().strip() for i in range(pages)]
    d.close()
    print(f"[4] 页数={pages}")
    for i, t in enumerate(texts, 1):
        print(f"    第{i}页提取文本: {t[:80]!r}")
    assert pages == 3, "页数应保持一致"
    assert any(t for t in texts), "文本层为空，OCR 失败"

    # 关键字搜索验证（模拟 Adobe Reader 的查找）
    d = fitz.open(str(out_pdf))
    joined = "\n".join(d[i].get_text() for i in range(d.page_count))
    d.close()
    kw = "OCR"
    print(f"[5] 全文搜索关键字 {kw!r}: {'命中' if kw in joined else '未命中'}")
    assert kw in joined, "文本层不可搜索"

    rep = WORK / "out" / mod.REPORT_FILENAME
    assert rep.exists(), f"未生成报告: {rep}"
    from openpyxl import load_workbook
    ws = load_workbook(rep).active
    print(f"[7] 报告: {rep}")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        if any(v is not None for v in row):
            print("    ", [str(v)[:28] for v in row if v is not None])

    print("\n=== 全部检查通过 ✅ ===")
    return 0


if __name__ == "__main__":
    sys.exit(main())
